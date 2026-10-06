import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from analytics_agent.models import TABLE_IDS
from analytics_agent.schema import (
    ColumnSchema,
    SchemaCatalog,
    TableSchema,
    load_dictionary,
    merge_bigquery_schema,
)


class DictionaryTests(unittest.TestCase):
    def make_workbook(self, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "th_primas_final"
        for row in rows:
            sheet.append(row)
        for name in ("mm_final", "mm_final_netos"):
            workbook.create_sheet(name)
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        handle.close()
        workbook.save(handle.name)
        return Path(handle.name)

    def test_loads_partial_dictionary_and_ignores_extra_exists_column(self):
        path = self.make_workbook([
            ("field name", "type", "description", "existe"),
            ("vrprima", "FLOAT", "Prima emitida", 0),
            ("nombre_ramo_comercial", "STRING", "Ramo", None),
        ])
        try:
            catalog = load_dictionary(path)
        finally:
            path.unlink(missing_ok=True)
        table = catalog.tables[TABLE_IDS["th_primas_final"].lower()]
        self.assertEqual({"vrprima", "nombre_ramo_comercial"}, set(table.columns))
        self.assertFalse(table.metadata_loaded)
        self.assertNotIn(TABLE_IDS["mm_final"].lower(), catalog.tables)

    def test_actual_bigquery_schema_adds_undocumented_fields(self):
        table_id = TABLE_IDS["th_primas_final"]
        catalog = SchemaCatalog(tables={table_id.lower(): TableSchema(
            "th_primas_final",
            table_id,
            {"vrprima": ColumnSchema("vrprima", "FLOAT", "Prima de Excel")},
        )})
        report = {
            table_id: {
                "ok": True,
                "actual_schema": {
                    "vrprima": {"name": "vrprima", "data_type": "NUMERIC", "description": "Prima BQ"},
                    "gross_anterior": {"name": "gross_anterior", "data_type": "FLOAT", "description": ""},
                },
                "location": "us-east1",
                "table_type": "TABLE",
            }
        }
        merge_bigquery_schema(catalog, report)
        table = catalog.tables[table_id.lower()]
        self.assertEqual({"vrprima", "gross_anterior"}, set(table.columns))
        self.assertEqual("Prima de Excel", table.columns["vrprima"].description)
        self.assertTrue(table.metadata_loaded)

    def test_search_supports_terms_and_pagination(self):
        table_ids = list(TABLE_IDS.values())
        catalog = SchemaCatalog(tables={
            table_id.lower(): TableSchema("", table_id, {}, description=description)
            for table_id, description in zip(table_ids, ("Primas emitidas", "Pagos de siniestros", "Pagos netos"))
        })
        matches, total = catalog.search_tables("primas", page_size=1)
        self.assertEqual(1, total)
        self.assertEqual(table_ids[0], matches[0].table_id)
        all_matches, all_count = catalog.search_tables("", offset=1, page_size=1)
        self.assertEqual(3, all_count)
        self.assertEqual(1, len(all_matches))

    def test_prompt_keeps_all_column_names_but_prioritizes_relevant_descriptions(self):
        table = TableSchema(
            sheet_name="demo",
            table_id="project.dataset.table",
            columns={
                "fecha_emision": ColumnSchema("fecha_emision", "DATE", "Fecha de emisión de la póliza."),
                "vrprima": ColumnSchema("vrprima", "NUMERIC", "Valor de prima emitida."),
                "documento": ColumnSchema("documento", "STRING", "Documento de identidad del cliente."),
            },
        )
        prompt = SchemaCatalog(tables={table.table_id: table}).prompt_text(query="prima emitida")
        self.assertIn("fecha_emision (DATE)", prompt)
        self.assertIn("vrprima (NUMERIC) — Valor de prima emitida.", prompt)
        self.assertIn("documento (STRING)", prompt)
        self.assertNotIn("Documento de identidad del cliente", prompt)

    def test_missing_workbook_tabs_do_not_disable_catalog(self):
        workbook = Workbook()
        workbook.active.title = "metadata"
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        handle.close()
        workbook.save(handle.name)
        try:
            catalog = load_dictionary(handle.name)
        finally:
            Path(handle.name).unlink(missing_ok=True)
        self.assertFalse(catalog.tables)
        self.assertTrue(catalog.warnings)


if __name__ == "__main__":
    unittest.main()
