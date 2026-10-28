"""Conversation regressions using actual compiled SQL and an independent SQLite oracle.

The fixture UDF uses a documented daily allocation for these tests only; its
signature is deliberately reversed. Production always uses the BigQuery UDF.
"""
import copy
import re
import sqlite3
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

import sqlglot
from openpyxl import Workbook, load_workbook

from analytics_agent.agent import AnalyticsAgent
from analytics_agent.analysis_context import AnalysisSpec, period_from_text, requested_metrics, resolve_analysis
from analytics_agent.analysis_sql import ROUTINE_ID, compile_analysis, routine_roles
from analytics_agent.bigquery_adapter import BigQueryError, DryRunResult, QueryPage
from analytics_agent.llm import Completion, OpenAICompatibleProvider
from analytics_agent.memory import SessionStore
from analytics_agent.models import AgentSettings
from analytics_agent.organization import HEADERS, OrganizationCatalog, import_workbook
from analytics_agent.presentation import analysis_presentation, number
from analytics_agent.schema import ColumnSchema, SchemaCatalog, TableSchema
from analytics_agent.semantics import MetricDefinitionError, load_metric_definitions
from analytics_agent.teams_common import result_page
from analytics_agent.teams_common import answer_record
from analytics_agent.cli import _format_cell
from tests.test_agent import action

PREMIUMS = "centralizacion-datos.analytics_curated.th_primas_final"
CLAIMS = "centralizacion-datos.analytics_curated.mm_final"
TODAY = date(2026, 10, 28)
SIGNATURE = {"routine_id": ROUTINE_ID, "routine_type": "SCALAR_FUNCTION", "language": "SQL", "remote": False,
             "arguments": [{"name": name, "data_type": "FLOAT64" if name == "primaemitida" else "DATE"}
                           for name in ("finvigencia", "iniciovigencia", "inicioanalisis", "finanalisis", "primaemitida")]}


class Oracle:
    location = "us-east1"

    def __init__(self, *, routine_available=True, fail_routine_sql=False):
        self.db = sqlite3.connect(":memory:")
        self.db.row_factory = sqlite3.Row
        self.db.execute("CREATE TABLE premiums(fecha_emision TEXT, fecha_iniciovig TEXT, fecha_finvig TEXT, vrprima REAL, cod_subprofit_loc INTEGER, linea_mayor TEXT, nombre_ramo_comercial TEXT)")
        self.db.executemany("INSERT INTO premiums VALUES(?,?,?,?,?,?,?)", [
            ("2025-12-15", "2026-01-01", "2026-12-31", 365, 34, "Copropiedades", "COPRO LEGADO"),
            ("2026-09-05", "2026-01-01", "2026-12-31", 730, 34, "Copropiedades", "COPRO NUEVO"),
            ("2026-09-05", "2026-01-01", "2026-12-31", 999, 51, "Otro", "COPROPIEDADES"),
        ])
        self.db.execute("CREATE TABLE claims(tipo TEXT, clase_pago TEXT, incurrido_total REAL, fecha_accidente TEXT, fecha_pago TEXT, mes_mmalla TEXT, lob_l1 TEXT, lob_l2 TEXT, lob_l3 TEXT, ramo_comercial TEXT)")
        self.db.executemany("INSERT INTO claims VALUES(?,?,?,?,?,?,?,?,?,?)", [
            ("Pagos", "Indemnización", 93, "2026-09-09", "2026-10-10", "2026-10-01", "Generales", "Copropiedades", "Copropiedades", "COPRO NUEVO"),
            ("Pagos", "Honorarios", -3, "2026-09-12", "2026-09-13", "2026-09-01", "Generales", "Copropiedades", "Copropiedades", "COPRO LEGADO"),
            ("Pagos", "Indemnización", 300, "2026-10-10", "2026-09-05", "2026-09-01", "Generales", "Copropiedades", "Copropiedades", "COPRO NUEVO"),
            ("Pagos", "Indemnización", 9999, "2026-09-09", "2026-09-09", "2026-09-01", "Generales", "Otro", "Otro", "COPROPIEDADES"),
            ("Avisado", None, 200, None, None, "2026-09-01", "Generales", "Copropiedades", "Copropiedades", "COPRO NUEVO"),
        ])
        def earned(end, start, first, last, premium):
            if None in (end, start, first, last, premium):
                return None
            end, start, first, last = (date.fromisoformat(value) for value in (end, start, first, last))
            overlap = max(0, (min(end, last) - max(start, first)).days + 1)
            return premium * overlap / ((end - start).days + 1)
        self.db.create_function("earned", 5, earned)
        self.executed = []
        self.routine_calls = 0
        self.routine_available = routine_available
        self.fail_routine_sql = fail_routine_sql
        self.columns = {
            PREMIUMS: {"fecha_emision": "DATE", "fecha_iniciovig": "DATE", "fecha_finvig": "DATE", "vrprima": "FLOAT64", "cod_subprofit_loc": "INT64", "linea_mayor": "STRING", "nombre_ramo_comercial": "STRING"},
            CLAIMS: {"tipo": "STRING", "clase_pago": "STRING", "incurrido_total": "FLOAT64", "fecha_accidente": "DATE", "fecha_pago": "DATE", "mes_mmalla": "DATE", "lob_l1": "STRING", "lob_l2": "STRING", "lob_l3": "STRING", "ramo_comercial": "STRING"},
        }

    def tables(self):
        return {key: TableSchema("", key, {k: ColumnSchema(k, v) for k, v in columns.items()}, metadata_loaded=True)
                for key, columns in self.columns.items()}

    def inspect_table_metadata(self, table_id):
        return {"table_id": table_id, "actual_schema": {k: {"name": k, "data_type": v} for k, v in self.columns[table_id].items()},
                "location": self.location, "table_type": "TABLE"}

    def inspect_routine_metadata(self, routine_id):
        self.routine_calls += 1
        if not self.routine_available:
            raise BigQueryError("Signature unavailable")
        return copy.deepcopy(SIGNATURE)

    def query(self, sql, parameters):
        local = sql.replace(f"`{PREMIUMS}`", "premiums").replace(f"`{CLAIMS}`", "claims").replace(f"`{ROUTINE_ID}`", "earned")
        bindings = {}
        for param in parameters:
            if param.get("array"):
                names = [param["name"] + "_" + str(i) for i in range(len(param["value"]))]
                local = local.replace("UNNEST(@" + param["name"] + ")", "(" + ",".join("@" + name for name in names) + ")")
                bindings.update(zip(names, param["value"]))
            else:
                bindings[param["name"]] = param["value"]
        local = sqlglot.transpile(local, read="bigquery", write="sqlite")[0]
        local = re.sub(r"DATE_SUB\(([^,]+), '1', DAY\)", r"DATE(\1, '-1 day')", local)
        return self.db.execute(local, bindings)

    def dry_run(self, sql, *, location="", parameters=()):
        if "INFORMATION_SCHEMA.PARAMETERS" in sql:
            return DryRunResult(self.routine_available, 0, "No signature access")
        if self.fail_routine_sql and "prima_devengada`(" in sql:
            return DryRunResult(False, 0, "Access Denied: routine prima_devengada")
        try:
            self.query(sql, parameters).fetchall()
            return DryRunResult(True, 100)
        except sqlite3.Error as error:
            return DryRunResult(False, 0, str(error))

    def execute(self, sql, *, location="", parameters=()):
        self.executed.append((sql, parameters))
        if "INFORMATION_SCHEMA.PARAMETERS" in sql:
            rows = tuple({"parameter_name": arg["name"], "data_type": arg["data_type"], "ordinal_position": i + 1,
                          "routine_type": "FUNCTION", "routine_body": "SQL", "external_language": None}
                         for i, arg in enumerate(SIGNATURE["arguments"]))
            columns = ("parameter_name", "data_type", "ordinal_position", "routine_type", "routine_body", "external_language")
        else:
            cursor = self.query(sql, parameters)
            columns = tuple(col[0] for col in cursor.description)
            rows = tuple(dict(row) for row in cursor.fetchall())
        return QueryPage("analysis-" + str(len(self.executed)), columns, rows, False, len(rows), 100)


class DirectLLM:
    def __init__(self, actions):
        self.actions = list(actions)
        self.calls = []

    def next_action(self, question, **kwargs):
        self.calls.append((question, kwargs))
        if not self.actions:
            raise AssertionError("An unnecessary additional model call was made")
        return self.actions.pop(0), {"total_tokens": 10}


class AnalysisTests(unittest.TestCase):
    def setUp(self):
        self.oracle = Oracle()
        self.addCleanup(self.oracle.db.close)
        self.rules = load_metric_definitions()

    def compile(self, question, action_data=None):
        spec = resolve_analysis(action_data or {}, question, today=TODAY)
        return compile_analysis(spec, self.rules, tables=self.oracle.tables(), timezone="America/Bogota", routine_metadata=SIGNATURE)

    def run_compiled(self, compiled):
        return [dict(row) for row in self.oracle.query(compiled.query.sql, compiled.query.parameters).fetchall()]

    def agent(self, actions, **settings):
        llm = DirectLLM(actions)
        return AnalyticsAgent(catalog=SchemaCatalog({}), llm=llm, bigquery=self.oracle,
                              settings=AgentSettings("https://example", "model", "secret", bigquery_location="us-east1", **settings)), llm

    def test_independent_sources_earned_includes_policies_issued_before_period(self):
        compiled = self.compile("Pagos, prima emitida, prima devengada y siniestralidad de Copropiedades en septiembre de 2026")
        rows = self.run_compiled(compiled)
        self.assertEqual(1, len(rows))
        self.assertEqual(730, rows[0]["prima_emitida"])
        self.assertEqual(90, rows[0]["prima_devengada"])
        self.assertEqual(90, rows[0]["pagos_siniestros"])
        self.assertEqual(100, rows[0]["siniestralidad_pagos"])
        # Two premium facts and two payment facts would multiply if joined raw.
        self.assertNotIn("JOIN `" + CLAIMS, compiled.query.sql)

    def test_last_six_months_are_exactly_six_and_current_is_partial(self):
        compiled = self.compile("Pagos, prima emitida, prima devengada y siniestralidad de Copropiedades de los últimos seis meses")
        rows = self.run_compiled(compiled)
        self.assertEqual(["2026-05", "2026-06", "2026-07", "2026-08", "2026-09", "2026-10"], [row["mes"] for row in rows])
        self.assertEqual([93, 90, 93, 93, 90, 84], [row["prima_devengada"] for row in rows])
        self.assertEqual(730, rows[-2]["prima_emitida"])
        self.assertEqual(300, rows[-1]["pagos_siniestros"])
        self.assertTrue(any("parcial" in note for note in compiled.spec.assumptions))

    def test_payment_basis_changes_results_and_reserves_keep_month_basis(self):
        accident = self.compile("Pagos de Copropiedades en septiembre de 2026")
        paid = self.compile("Pagos de Copropiedades en septiembre de 2026 usando fecha_pago")
        self.assertEqual(90, self.run_compiled(accident)[0]["pagos_siniestros"])
        self.assertEqual(297, self.run_compiled(paid)[0]["pagos_siniestros"])
        reserves = self.compile("Reservas de Copropiedades en septiembre de 2026 por fecha_pago")
        self.assertEqual(200, self.run_compiled(reserves)[0]["reservas_avisadas"])
        self.assertIn("mes_mmalla", reserves.query.sql)

    def test_honorarios_and_indemnizaciones_use_same_basis_and_class_cache(self):
        agent, _ = self.agent([action("run_analysis"), action("run_analysis")])
        first = agent.answer("Honorarios e indemnizaciones de Copropiedades septiembre de 2026")
        self.assertEqual(-3, first.summary_table.rows[0]["honorarios"])
        self.assertEqual(93, first.summary_table.rows[0]["indemnizaciones"])
        second = agent.answer("Honorarios e indemnizaciones de Copropiedades septiembre de 2026 por fecha_pago")
        self.assertEqual(300, second.summary_table.rows[0]["indemnizaciones"])
        self.assertEqual(1, sum("SELECT DISTINCT" in sql for sql, _ in self.oracle.executed))

    def test_area_mapping_and_exact_organizational_population(self):
        compiled = self.compile("Primas emitidas y pagos por área en septiembre de 2026")
        rows = self.run_compiled(compiled)
        self.assertEqual(2, len(rows))
        coprop = next(row for row in rows if row["linea_mayor"] == "Copropiedades")
        self.assertEqual(730, coprop["prima_emitida"])
        self.assertEqual(90, coprop["pagos_siniestros"])
        self.assertIn("`lob_l2` AS `linea_mayor`", compiled.query.sql)
        scoped = self.compile("Primas emitidas de Copropiedades en septiembre de 2026")
        self.assertEqual([34], scoped.spec.population["codes"])
        self.assertNotIn("nombre_ramo_comercial", scoped.query.sql)

    def test_followups_preserve_scope_filters_grain_and_only_change_requested_fields(self):
        initial = resolve_analysis({"group_by": ["area"], "filters": [{"field": "canal", "operator": "eq", "value": "Banco"}]},
                                   "Pagos y prima emitida de Copropiedades últimos seis meses", today=TODAY)
        for question in ("Vuelve a intentarlo", "Confirma esos valores", "Usa fecha_pago"):
            changed = resolve_analysis({"group_by": [], "filters": [], "time_grain": "total", "population": {"label": "PYMES"}},
                                       question, today=TODAY + timedelta(days=2), active=initial.to_dict())
            self.assertEqual(initial.population, changed.population)
            self.assertEqual(initial.start_date, changed.start_date)
            self.assertEqual(initial.end_date, changed.end_date)
            self.assertEqual(initial.metric_keys, changed.metric_keys)
            self.assertEqual(initial.filters, changed.filters)
            self.assertEqual(initial.group_by, changed.group_by)
            self.assertEqual("month", changed.time_grain)
        self.assertEqual("payment", changed.payment_date_basis)
        fresh = resolve_analysis({}, "Primas emitidas en septiembre de 2026", today=TODAY, active=initial.to_dict())
        self.assertFalse(fresh.population)
        self.assertFalse(fresh.filters)
        cleared = resolve_analysis({}, "Ahora total general sin filtros", today=TODAY, active=initial.to_dict())
        self.assertEqual("total", cleared.time_grain)
        self.assertFalse(cleared.population)
        self.assertFalse(cleared.filters)
        self.assertFalse(cleared.group_by)

    def test_no_invented_population_or_other_premium_definition(self):
        spec = resolve_analysis({"metric_keys": ["prima_emitida"], "population": {"label": "Copropiedades"}}, "Primas de septiembre de 2026", today=TODAY)
        self.assertFalse(spec.population)
        self.assertEqual((), requested_metrics("Primas retenidas netas de septiembre"))
        self.assertEqual(("pagos_siniestros",), requested_metrics("Todos los pagos, incluyendo honorarios"))
        with self.assertRaises(MetricDefinitionError):
            resolve_analysis({"metric_keys": ["prima_emitida"]}, "Primas retenidas de septiembre", today=TODAY)

    def test_large_unbounded_new_analysis_does_not_borrow_old_dates_to_skip_question(self):
        with tempfile.TemporaryDirectory() as folder, SessionStore(folder) as store:
            session = store.create_session()
            turn = store.begin_turn(session, "Pagos de septiembre de 2026")
            store.complete_turn(turn, "Resultado anterior")
            agent, _ = self.agent([action("run_analysis")], unscoped_query_threshold_bytes=1)
            result = agent.answer("Primas emitidas", session_store=store, session_id=session)
            self.assertEqual("date_scope", result.clarification.reason)
            self.assertEqual(0, len(self.oracle.executed))

    def test_explicit_branch_and_organization_are_distinct_scopes(self):
        compiled = self.compile("Primas emitidas del ramo comercial COPROPIEDADES en septiembre de 2026")
        self.assertFalse(compiled.spec.population)
        self.assertEqual(999, self.run_compiled(compiled)[0]["prima_emitida"])
        organizational = self.compile("Primas emitidas de Copropiedades en septiembre de 2026")
        grouped = resolve_analysis({}, "Ahora por ramo", today=TODAY, active=organizational.spec.to_dict())
        self.assertEqual([34], grouped.population["codes"])
        self.assertEqual(("ramo",), grouped.group_by)

    def test_followup_ratio_changes_denominator_only(self):
        initial = self.compile("Siniestralidad de Copropiedades en septiembre de 2026").spec
        changed = resolve_analysis({}, "Usa prima emitida", today=TODAY, active=initial.to_dict())
        self.assertEqual(("siniestralidad_pagos_emitida",), changed.metric_keys)
        self.assertEqual(initial.population, changed.population)
        self.assertEqual(initial.start_date, changed.start_date)

    def test_signature_preserves_ordinal_order_and_rejects_unknown_remote_and_types(self):
        self.assertEqual(("end_coverage", "start_coverage", "start_analysis", "end_analysis", "premium"), routine_roles(SIGNATURE))
        for mutate in (lambda sig: sig.update(remote=True),
                       lambda sig: sig["arguments"][0].update(name="fecha"),
                       lambda sig: sig["arguments"][0].update(data_type="STRING"),
                       lambda sig: sig["arguments"].pop()):
            sig = copy.deepcopy(SIGNATURE)
            mutate(sig)
            with self.assertRaises(MetricDefinitionError):
                routine_roles(sig)

    def test_null_zero_no_movements_and_zero_denominator_are_distinct(self):
        compiled = self.compile("Siniestralidad de Copropiedades septiembre de 2026")
        row = self.run_compiled(compiled)[0]
        row.update(prima_devengada=0, siniestralidad_pagos=None)
        _, table, notes = analysis_presentation({"columns": list(row), "rows": [row]}, compiled)
        self.assertEqual("No calculable", number(table.rows[0]["siniestralidad_pagos"], percent=True))
        self.assertTrue(any("cero" in note for note in notes))
        row.update(prima_devengada=None, __rows_prima_devengada=1, __values_prima_devengada=0)
        _, _, notes = analysis_presentation({"columns": list(row), "rows": [row]}, compiled)
        self.assertTrue(any("sin importes" in note for note in notes))
        row.update(__rows_prima_devengada=0)
        _, _, notes = analysis_presentation({"columns": list(row), "rows": [row]}, compiled)
        self.assertTrue(any("no se encontraron movimientos" in note for note in notes))
        self.assertEqual("0,00", number(0))

    def test_one_model_call_fast_answer_and_excel_history_learning_and_session_isolation(self):
        with tempfile.TemporaryDirectory() as folder, SessionStore(folder) as store:
            agent, llm = self.agent([action("run_analysis"), action("run_analysis")], export_dir=str(Path(folder) / "exports"))
            first = agent.answer("Pagos, prima emitida y siniestralidad de Copropiedades septiembre de 2026 en Excel", session_store=store)
            self.assertEqual(1, len(llm.calls))
            self.assertEqual(1, first.query_count)
            self.assertEqual(100, first.summary_table.rows[0]["siniestralidad_pagos"])
            self.assertTrue(any(param["value"] == [34] for param in first.audit[-1]["parameters"]))
            self.assertIn("100,00 %", first.answer)
            self.assertEqual(1, len(first.artifacts))
            wb = load_workbook(first.artifacts[0]["path"], read_only=True)
            self.addCleanup(wb.close)
            self.assertTrue(any("Prima emitida" in row for sheet in wb for row in sheet.iter_rows(values_only=True)))
            second = agent.answer("Usa fecha_pago", session_store=store, session_id=first.session_id)
            self.assertEqual(297, second.summary_table.rows[0]["pagos_siniestros"])
            self.assertEqual(330, second.summary_table.rows[0]["siniestralidad_pagos"])
            self.assertEqual(1, self.oracle.routine_calls)
            self.assertEqual(2, len(llm.calls))
            active = store.get_analysis_state(first.session_id)
            self.assertEqual([34], active["population"]["codes"])
            self.assertEqual("2026-09-01", active["start_date"])
            session = store.create_session()
            self.assertEqual({}, store.get_analysis_state(session))
            self.assertEqual(2, store.connection.execute("SELECT COUNT(*) FROM learning_queue WHERE status='pending'").fetchone()[0])
            saved = store.recent_query_results(first.session_id)[0]
            self.assertEqual("answer", saved["metadata"]["result_role"])

    def test_missing_signature_keeps_available_values_and_marks_history_partial(self):
        self.oracle.routine_available = False
        agent, llm = self.agent([action("run_analysis")])
        with tempfile.TemporaryDirectory() as folder, SessionStore(folder) as store:
            result = agent.answer("Pagos, prima emitida y siniestralidad de Copropiedades septiembre de 2026", session_store=store)
            self.assertEqual("partial", result.status)
            self.assertEqual(730, result.summary_table.rows[0]["prima_emitida"])
            self.assertEqual(90, result.summary_table.rows[0]["pagos_siniestros"])
            self.assertIsNone(result.summary_table.rows[0]["siniestralidad_pagos"])
            self.assertIn("Resultado parcial", result.answer)
            self.assertEqual("partial", store.connection.execute("SELECT status FROM turns WHERE turn_id=?", (result.turn_id,)).fetchone()[0])

    def test_udf_execution_denied_preserves_components(self):
        self.oracle.fail_routine_sql = True
        agent, _ = self.agent([action("run_analysis")])
        result = agent.answer("Pagos, prima emitida y siniestralidad de Copropiedades septiembre de 2026")
        self.assertEqual("partial", result.status)
        self.assertEqual(90, result.summary_table.rows[0]["pagos_siniestros"])
        self.assertEqual(730, result.summary_table.rows[0]["prima_emitida"])

    def test_missing_coverage_field_does_not_erase_other_metrics(self):
        self.oracle.columns[PREMIUMS].pop("fecha_iniciovig")
        agent, _ = self.agent([action("run_analysis")])
        result = agent.answer("Pagos, prima emitida y siniestralidad de Copropiedades septiembre de 2026")
        self.assertEqual("partial", result.status)
        self.assertEqual(730, result.summary_table.rows[0]["prima_emitida"])
        self.assertEqual(90, result.summary_table.rows[0]["pagos_siniestros"])
        self.assertIsNone(result.summary_table.rows[0]["prima_devengada"])

    def test_routine_information_schema_fallback(self):
        agent, _ = self.agent([action("run_analysis")])
        with patch.object(self.oracle, "inspect_routine_metadata", side_effect=BigQueryError("metadata unavailable")):
            result = agent.answer("Siniestralidad de Copropiedades septiembre de 2026")
        self.assertEqual("complete", result.status)
        self.assertEqual(2, result.query_count)
        self.assertEqual(100, result.summary_table.rows[0]["siniestralidad_pagos"])

    def test_remote_signature_is_not_reclassified_as_sql_by_fallback(self):
        agent, _ = self.agent([action("run_analysis")])
        signature = copy.deepcopy(SIGNATURE)
        signature["remote"] = True
        with patch.object(self.oracle, "inspect_routine_metadata", return_value=signature):
            result = agent.answer("Pagos y siniestralidad de Copropiedades septiembre de 2026")
        self.assertEqual("partial", result.status)
        self.assertEqual(1, len(self.oracle.executed))
        self.assertNotIn("INFORMATION_SCHEMA", self.oracle.executed[0][0])

    def test_routine_cache_expiry_and_percent_rendering_in_both_channels(self):
        agent, _ = self.agent([action("run_analysis"), action("run_analysis")])
        first = agent.answer("Siniestralidad de Copropiedades septiembre de 2026")
        self.assertEqual(1, self.oracle.routine_calls)
        agent._routine_cache[ROUTINE_ID] = (-4000, SIGNATURE)
        agent.answer("Siniestralidad de Copropiedades septiembre de 2026")
        self.assertEqual(2, self.oracle.routine_calls)
        page = result_page(answer_record(first), 0)
        self.assertEqual("100,00 %", page["rows"][0]["siniestralidad_pagos"])
        self.assertEqual("Siniestralidad sobre devengada", page["column_labels"]["siniestralidad_pagos"])
        self.assertEqual("100,00 %", _format_cell(100, format_kind="percent"))

    def test_diagnostic_table_cannot_replace_primary_summary(self):
        results = {"amount": {"result_role": "answer", "columns": ["prima_emitida"], "rows": [{"prima_emitida": 100}]},
                   "count": {"result_role": "diagnostic", "columns": ["filas"], "rows": [{"filas": 10}]}}
        agent, _ = self.agent([])
        summary = agent._summary_from_evidence(None, results, {"amount", "count"})
        self.assertEqual(("prima_emitida",), summary.columns)

    def test_repeated_inspections_do_not_start_another_query_or_model_loop(self):
        agent, llm = self.agent([
            action("run_select", sql=f"SELECT SUM(vrprima) AS prima_emitida FROM `{PREMIUMS}` WHERE fecha_emision >= DATE '2026-09-01'"),
            action("inspect_result", query_id="analysis-1", operation="profile", column="", offset=0, limit=20),
            action("inspect_result", query_id="analysis-1", operation="profile", column="", offset=0, limit=20, notes="different wording"),
        ])
        result = agent.answer("Prima emitida de septiembre de 2026")
        self.assertEqual(1, result.query_count)
        self.assertEqual(3, len(llm.calls))
        self.assertEqual(1729, result.summary_table.rows[0]["prima_emitida"])

    def test_period_variants(self):
        for text, start, end in (("septiembre", "2026-09-01", "2026-10-01"),
                                 ("ayer", "2026-10-27", "2026-10-28"),
                                 ("año pasado", "2025-01-01", "2026-01-01"),
                                 ("primer semestre de 2026", "2026-01-01", "2026-07-01"),
                                 ("diciembre de 2025 a enero de 2026", "2025-12-01", "2026-02-01"),
                                 ("del 1 de enero al 30 de junio de 2026", "2026-01-01", "2026-07-01"),
                                 ("últimos seis meses cerrados", "2026-04-01", "2026-10-01")):
            self.assertEqual((start, end), period_from_text(text, today=TODAY)[:2])

    def test_action_parser_leaves_omitted_followup_fields_unset(self):
        class Provider(OpenAICompatibleProvider):
            def complete(self, messages, *, max_tokens=1200):
                return Completion('{"action":"run_analysis","metric_keys":[]}', {})
        provider = Provider(base_url="https://example", model="m", api_key="test")
        parsed, _ = provider.next_action("Usa fecha_pago")
        for key in ("group_by", "filters", "time_grain", "population"):
            self.assertNotIn(key, parsed)


class OrganizationTests(unittest.TestCase):
    def test_catalog_unique_codes_and_separate_many_to_one_bridge(self):
        catalog = OrganizationCatalog()
        self.assertEqual(98, catalog.payload["source_rows"])
        self.assertEqual(82, len(catalog.rows))
        self.assertEqual(98, len(catalog.payload["subprofit_links"]))
        self.assertEqual([34], catalog.resolve("Copropiedades")["codes"])
        self.assertEqual("PYMES", catalog.resolve("Copropiedades")["paths"][0]["nombre_lob_mor"])
        self.assertTrue(any(link == {"cod_subprofit": 51, "cod_subprofit_loc": 34} for link in catalog.payload["subprofit_links"]))
        self.assertIsNone(catalog.from_question("Primas del ramo comercial COPROPIEDADES"))

    def test_importer_deduplicates_and_rejects_conflicting_hierarchies(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "hierarchy.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(HEADERS)
            sheet.append(["Generales", "Copropiedades", "Copropiedades", 51, 34, "PYMES"])
            sheet.append(["Generales", "Copropiedades", "Copropiedades", 52, 34, "PYMES"])
            workbook.save(path)
            payload = import_workbook(path)
            self.assertEqual(1, len(payload["local_codes"]))
            self.assertEqual(2, len(payload["subprofit_links"]))
            sheet.append(["Generales", "Otro", "Otro", 53, 34, "PYMES"])
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "contradictorias"):
                import_workbook(path)
            workbook.close()

    def test_graph_connects_codes_hierarchy_metric_and_function(self):
        with tempfile.TemporaryDirectory() as folder, SessionStore(folder) as store:
            store.knowledge("Copropiedades")
            self.assertEqual(82, store.connection.execute("SELECT COUNT(*) FROM knowledge_nodes WHERE kind='local_code'").fetchone()[0])
            edge = store.connection.execute("SELECT target FROM knowledge_edges WHERE source='local_code:34' AND relation='belongs_to_linea_mayor'").fetchone()[0]
            self.assertEqual("organization:linea_mayor:Copropiedades", edge)
            self.assertEqual([34], store.knowledge("Copropiedades")["organization"][0]["codes"])
            self.assertTrue(store.connection.execute("SELECT 1 FROM knowledge_edges WHERE source='metric:prima_devengada' AND relation='uses_routine'").fetchone())


if __name__ == "__main__":
    unittest.main()
