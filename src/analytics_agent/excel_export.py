from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Mapping, Sequence
from uuid import uuid4
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
_NUMERIC_TYPES = {"INT64", "INTEGER", "FLOAT64", "FLOAT", "NUMERIC", "BIGNUMERIC", "DECIMAL", "BIGDECIMAL"}


def _safe_sheet_name(value: str, existing: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", value).strip(" '") or "Datos"
    base = base[:31]
    candidate = base
    suffix = 2
    while candidate.lower() in {name.lower() for name in existing}:
        tail = f"_{suffix}"
        candidate = base[:31 - len(tail)] + tail
        suffix += 1
    existing.add(candidate)
    return candidate


def _cell_value(worksheet: Any, value: Any, data_type: str = "") -> WriteOnlyCell:
    """Build a literal Excel cell and preserve BigQuery numeric precision."""
    type_parts = data_type.upper().split()
    kind = type_parts[0] if type_parts else ""
    if value is None:
        return WriteOnlyCell(worksheet, value=None)
    if kind in _NUMERIC_TYPES and not isinstance(value, bool):
        try:
            if kind in {"INT64", "INTEGER"}:
                numeric = int(str(value))
                if len(str(abs(numeric))) <= 15:
                    return WriteOnlyCell(worksheet, value=numeric)
            elif kind in {"NUMERIC", "BIGNUMERIC", "DECIMAL", "BIGDECIMAL"}:
                numeric = Decimal(str(value))
                digits = len(numeric.as_tuple().digits)
                if numeric.is_finite() and digits <= 15:
                    return WriteOnlyCell(worksheet, value=numeric)
            elif kind in {"FLOAT64", "FLOAT"}:
                number = float(value)
                if number == number and abs(number) != float("inf"):
                    return WriteOnlyCell(worksheet, value=number)
        except (ValueError, TypeError, InvalidOperation, OverflowError):
            pass

    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    cell = WriteOnlyCell(worksheet, value=str(value) if not isinstance(value, (int, float, bool)) else value)
    # Force untrusted text to remain text, including values beginning with '='.
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def export_results_to_excel(
    results: Sequence[Mapping[str, Any]],
    *,
    title: str,
    export_dir: str | Path,
    timezone_name: str = "America/Bogota",
) -> Path:
    """Write one or more previously-read query results to a local .xlsx file."""
    if not results:
        raise ValueError("No hay resultados seleccionados para exportar")
    zone = ZoneInfo(timezone_name)
    exported_at = datetime.now(timezone.utc).astimezone(zone)
    directory = Path(export_dir).expanduser().resolve()
    directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    try:
        directory.chmod(0o700)
    except OSError:
        pass
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", title.strip()).strip("_")[:48] or "resultados"
    path = directory / f"{exported_at:%Y%m%d_%H%M%S}_{stem}_{uuid4().hex[:8]}.xlsx"

    workbook = Workbook(write_only=True)
    workbook.properties.title = title[:255]
    workbook.properties.creator = "Analytics Agent"
    summary = workbook.create_sheet("Resumen")
    header_fill = PatternFill("solid", fgColor="163A5F")
    header_font = Font(bold=True, color="FFFFFF")
    summary.append([_cell_value(summary, "Exportación"), _cell_value(summary, title[:255])])
    summary.append([_cell_value(summary, "Generado"), _cell_value(summary, exported_at.isoformat(timespec="seconds"))])
    summary.append([_cell_value(summary, "Zona horaria"), _cell_value(summary, timezone_name)])
    summary.append([])
    summary_header = [
        _cell_value(summary, label)
        for label in ("Consulta", "Filas", "Tablas", "ID del resultado", "Columnas", "SQL")
    ]
    for cell in summary_header:
        cell.fill = header_fill
        cell.font = header_font
    summary.append(summary_header)
    for column_index, width in enumerate((28, 14, 36, 36, 48, 72), start=1):
        summary.column_dimensions[get_column_letter(column_index)].width = width

    used_names = {"Resumen"}
    for result_number, result in enumerate(results, start=1):
        columns = [str(column) for column in result.get("columns", [])]
        rows = result.get("rows", [])
        if not columns:
            raise ValueError("El resultado seleccionado no contiene columnas")
        if len(columns) > EXCEL_MAX_COLUMNS:
            raise ValueError(f"El resultado tiene {len(columns)} columnas; Excel admite máximo {EXCEL_MAX_COLUMNS}")
        if not isinstance(rows, Sequence) or isinstance(rows, (str, bytes)):
            raise ValueError("Las filas del resultado no tienen un formato válido")
        query_label = str(result.get("title") or f"Resultado {result_number}")
        tables = result.get("tables", [])
        if isinstance(tables, str):
            tables = [tables]
        sql = str(result.get("sql", ""))
        summary.append([
            _cell_value(summary, query_label[:255]),
            _cell_value(summary, len(rows), "INT64"),
            _cell_value(summary, ", ".join(str(item) for item in tables)[:32000]),
            _cell_value(summary, str(result.get("query_id", ""))[:255]),
            _cell_value(summary, ", ".join(columns)[:32000]),
            _cell_value(summary, sql[:32000]),
        ])
        types = result.get("column_types", {})
        if not isinstance(types, Mapping):
            types = {}
        chunk_index = 1

        def new_data_sheet(index: int) -> Any:
            sheet_name = _safe_sheet_name(
                query_label if index == 1 else f"{query_label}_{index}", used_names
            )
            data_sheet = workbook.create_sheet(sheet_name)
            data_sheet.freeze_panes = "A2"
            data_rows = min(EXCEL_MAX_ROWS - 1, max(0, len(rows) - (index - 1) * (EXCEL_MAX_ROWS - 1)))
            data_sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{data_rows + 1}"
            header = []
            for column_index, column in enumerate(columns, start=1):
                cell = WriteOnlyCell(data_sheet, value=column)
                cell.data_type = "s"
                cell.fill = header_fill
                cell.font = header_font
                header.append(cell)
                data_sheet.column_dimensions[get_column_letter(column_index)].width = min(max(len(column) + 2, 12), 36)
            data_sheet.append(header)
            return data_sheet

        sheet = new_data_sheet(chunk_index)
        for row_index, row in enumerate(rows):
            if row_index and row_index % (EXCEL_MAX_ROWS - 1) == 0:
                chunk_index += 1
                sheet = new_data_sheet(chunk_index)
            if not isinstance(row, Mapping):
                continue
            sheet.append([_cell_value(sheet, row.get(column), str(types.get(column, ""))) for column in columns])

    workbook.save(path)
    try:
        path.chmod(0o600)
    except OSError:
        pass
    return path
