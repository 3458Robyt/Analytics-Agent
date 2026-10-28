"""Reviewed organizational taxonomy, separate from insurance branches and facts."""
from __future__ import annotations

import argparse
import json
import re
import unicodedata
from importlib.resources import files
from pathlib import Path
from typing import Any, Mapping

from .semantics import MetricDefinitionError


LEVELS = ("vicepresidencia", "linea_mayor", "linea_menor", "nombre_lob_mor")
HEADERS = (*LEVELS[:3], "cod_subprofit", "cod_subprofit_loc", "nombre_lob_mor")
DIMENSION_ALIASES = {
    "area": "linea_mayor", "areas": "linea_mayor", "vp": "vicepresidencia",
    "lob1": "vicepresidencia", "lob_l1": "vicepresidencia", "lob l1": "vicepresidencia",
    "lob2": "linea_mayor", "lob_l2": "linea_mayor", "lob l2": "linea_mayor",
    "lob3": "linea_menor", "lob_l3": "linea_menor", "lob l3": "linea_menor",
    "linea mayor": "linea_mayor", "linea menor": "linea_menor",
    "vicepresidencia": "vicepresidencia", "mor": "nombre_lob_mor",
}


def normalized(value: Any) -> str:
    return " ".join("".join(c for c in unicodedata.normalize("NFKD", str(value or "").lower())
                            if not unicodedata.combining(c)).split())


def dimension_name(value: str) -> str:
    return DIMENSION_ALIASES.get(normalized(value), value)


def import_workbook(path: str | Path) -> dict[str, Any]:
    """Build a unique local-code dimension and a separate many-to-one bridge."""
    from openpyxl import load_workbook

    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    dimension: dict[int, dict[str, Any]] = {}
    bridge: set[tuple[int, int]] = set()
    count = 0
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(str(v or "").strip() for v in next(rows))
        if not set(HEADERS).issubset(headers):
            raise ValueError("El Excel debe incluir: " + ", ".join(HEADERS))
        for number, row in enumerate(rows, 2):
            if not any(v is not None for v in row):
                continue
            record = dict(zip(headers, row))
            for field in LEVELS:
                if not isinstance(record[field], str) or not record[field].strip():
                    raise ValueError(f"Fila {number}: falta {field}")
            for field in ("cod_subprofit_loc", "cod_subprofit"):
                value = record[field]
                if isinstance(value, bool) or not re.fullmatch(r"\d+", str(value)):
                    raise ValueError(f"Fila {number}: {field} debe ser un código entero")
            code = int(record["cod_subprofit_loc"])
            item = {"cod_subprofit_loc": code, **{key: record[key].strip() for key in LEVELS}}
            if code in dimension and dimension[code] != item:
                raise ValueError(f"El código local {code} tiene clasificaciones contradictorias")
            dimension[code] = item
            bridge.add((int(record["cod_subprofit"]), code))
            count += 1
        return {
            "version": 1, "source": path.name, "reviewed_by": "David Centeno",
            "reviewed_on": "2026-10-28", "temporal_scope": "current_catalog_no_historical_validity",
            "source_rows": count,
            "local_codes": [dimension[key] for key in sorted(dimension)],
            "subprofit_links": [{"cod_subprofit": raw, "cod_subprofit_loc": local}
                                for raw, local in sorted(bridge)],
        }
    finally:
        workbook.close()


class OrganizationCatalog:
    def __init__(self, payload: Mapping[str, Any] | None = None) -> None:
        self.payload = dict(payload) if payload is not None else json.loads(
            files("analytics_agent").joinpath("data", "organization_catalog.json").read_text(encoding="utf-8"))
        self.rows = tuple(self.payload["local_codes"])
        codes = [row["cod_subprofit_loc"] for row in self.rows]
        if len(codes) != len(set(codes)):
            raise ValueError("El catálogo debe tener una sola fila por código local")

    def resolve(self, text: str, *, level: str = "") -> dict[str, Any] | None:
        """Exact normalized labels, never financial facts or arbitrary fuzzy matches."""
        value = normalized(text)
        level = dimension_name(level) if level else ""
        if level and level not in (*LEVELS, "cod_subprofit_loc"):
            raise MetricDefinitionError("Nivel organizacional desconocido")
        if level == "cod_subprofit_loc":
            matched = [row for row in self.rows if str(row["cod_subprofit_loc"]) == value]
        else:
            levels = (level,) if level else ("linea_mayor", "linea_menor", "vicepresidencia", "nombre_lob_mor")
            matched = []
            for candidate in levels:
                matched = [row for row in self.rows if normalized(row[candidate]) == value]
                if matched:
                    level = candidate
                    break
        if not matched:
            return None
        return {"level": level, "label": matched[0][level],
                "codes": sorted({row["cod_subprofit_loc"] for row in matched}),
                "paths": [dict(row) for row in matched], "source": self.payload["source"]}

    def from_question(self, question: str) -> dict[str, Any] | None:
        text = normalized(question)
        code = re.search(r"\b(?:cod_subprofit_loc|codigo local)\s*(?:=|es|:)?\s*(\d+)\b", text)
        if code:
            result = self.resolve(code[1], level="cod_subprofit_loc")
            if result is None:
                raise MetricDefinitionError(f"El código local {code[1]} no está en el catálogo revisado")
            return result
        # A branch explicitly named by the employee remains a branch filter.
        if re.search(r"\bramo(?:s)?(?: comercial(?:es)?)?\s+(?:de\s+)?[\w]", text):
            return None
        candidates = []
        for row in self.rows:
            for level in ("linea_mayor", "linea_menor", "vicepresidencia", "nombre_lob_mor"):
                label = normalized(row[level])
                if len(label) >= 4 and label not in {"generales", "personas"} and re.search(
                        r"(?<!\w)" + re.escape(label) + r"(?!\w)", text):
                    candidates.append((len(label), row[level]))
        # Generic VP labels need explicit organizational wording.
        for label in ("Generales", "Personas"):
            if re.search(r"\b(?:vp|vicepresidencia)\s+(?:de\s+)?" + normalized(label) + r"\b", text):
                candidates.append((len(label), label))
        return self.resolve(max(candidates, key=lambda item: item[0])[1]) if candidates else None

    def search(self, text: str) -> list[dict[str, Any]]:
        query = normalized(text)
        tokens = {t for t in re.findall(r"\w+", query) if len(t) > 2}
        result = self.from_question(text)
        if result:
            return [result]
        labels = {(level, row[level]) for row in self.rows for level in LEVELS}
        ranked = []
        for level, label in labels:
            score = sum(token in normalized(label) for token in tokens)
            if score:
                ranked.append((score, level, label))
        return [self.resolve(label, level=level) for _, level, label in sorted(ranked, reverse=True)[:4]]


def main() -> None:
    parser = argparse.ArgumentParser(description="Importar el catálogo organizacional revisado")
    parser.add_argument("workbook")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    payload = import_workbook(args.workbook)
    Path(args.output).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{payload['source_rows']} relaciones; {len(payload['local_codes'])} códigos locales únicos")


if __name__ == "__main__":
    main()
