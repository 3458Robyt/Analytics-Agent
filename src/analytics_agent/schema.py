from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping

from .models import TABLE_IDS


class DictionaryError(ValueError):
    """The optional workbook could not be read."""


def _normalize(value: Any) -> str:
    raw = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(c for c in raw if not unicodedata.combining(c)).strip().lower()


def _key(table_id: str) -> str:
    return table_id.strip().replace(":", ".").lower()


@dataclass(frozen=True)
class ColumnSchema:
    name: str
    data_type: str
    description: str = ""


@dataclass(frozen=True)
class TableSchema:
    sheet_name: str
    table_id: str
    columns: Mapping[str, ColumnSchema]
    warnings: tuple[str, ...] = field(default_factory=tuple)
    description: str = ""
    table_type: str = ""
    location: str = ""
    metadata_loaded: bool = False


@dataclass
class SchemaCatalog:
    """Workbook hints and live metadata for tables visible to the runtime."""

    tables: dict[str, TableSchema]
    warnings: tuple[str, ...] = field(default_factory=tuple)

    def add_table(self, table: TableSchema) -> None:
        key = _key(table.table_id)
        previous = self.tables.get(key)
        if previous:
            columns = dict(previous.columns)
            columns.update(table.columns)
            table = TableSchema(
                sheet_name=table.sheet_name or previous.sheet_name,
                table_id=table.table_id,
                columns=columns,
                warnings=tuple(dict.fromkeys((*previous.warnings, *table.warnings))),
                description=table.description or previous.description,
                table_type=table.table_type or previous.table_type,
                location=table.location or previous.location,
                metadata_loaded=table.metadata_loaded or previous.metadata_loaded,
            )
        self.tables[key] = table

    def prompt_text(self, table_ids: tuple[str, ...] | list[str] | None = None) -> str:
        selected = (
            [self.tables[_key(table_id)] for table_id in table_ids if _key(table_id) in self.tables]
            if table_ids is not None
            else list(self.tables.values())
        )
        sections: list[str] = []
        for table in selected:
            header = f"{table.table_id}"
            if table.table_type:
                header += f" [{table.table_type}]"
            if table.description:
                header += f" — {table.description}"
            fields = []
            for column in table.columns.values():
                description = f" — {column.description}" if column.description else ""
                fields.append(f"  - {column.name} ({column.data_type}){description}")
            sections.append(header + ("\n" + "\n".join(fields) if fields else "\n  - Esquema aún no consultado"))
        return "\n\n".join(sections)

    def search_tables(
        self,
        query: str,
        *,
        offset: int = 0,
        page_size: int = 100,
    ) -> tuple[list[TableSchema], int]:
        if offset < 0 or page_size < 1:
            raise ValueError("offset debe ser >= 0 y page_size debe ser > 0")
        terms = [_normalize(item) for item in re.findall(r"[\w.-]+", query) if item.strip()]
        matches: list[tuple[int, str, TableSchema]] = []
        for key, table in self.tables.items():
            haystack = _normalize(" ".join((table.table_id, table.description, *(c.name for c in table.columns.values()), *(c.description for c in table.columns.values()))))
            score = sum(3 if term in _normalize(table.table_id) else 1 for term in terms if term in haystack)
            if not terms or score:
                matches.append((score, key, table))
        matches.sort(key=lambda item: (-item[0], item[1]))
        total = len(matches)
        return [item[2] for item in matches[offset:offset + page_size]], total


def merge_bigquery_schema(
    catalog: SchemaCatalog,
    metadata_report: Mapping[str, Mapping[str, Any]],
) -> SchemaCatalog:
    """Merge available live fields with optional workbook descriptions."""
    for table_id, status in metadata_report.items():
        actual_schema = status.get("actual_schema")
        if not status.get("ok") or not isinstance(actual_schema, Mapping):
            continue
        existing = catalog.tables.get(_key(table_id))
        documented_columns = existing.columns if existing else {}
        merged_columns: dict[str, ColumnSchema] = {}
        for key, actual in actual_schema.items():
            if not isinstance(actual, Mapping):
                continue
            name = str(actual.get("name") or key).strip()
            normalized_name = name.lower()
            documented = documented_columns.get(normalized_name)
            description = (
                (documented.description.strip() if documented else "")
                or str(actual.get("description") or "").strip()
            )
            merged_columns[normalized_name] = ColumnSchema(
                name=name,
                data_type=str(actual.get("data_type") or "").strip().upper(),
                description=description,
            )
        catalog.add_table(TableSchema(
            sheet_name=existing.sheet_name if existing else "",
            table_id=str(status.get("table_id") or table_id),
            columns=merged_columns,
            warnings=existing.warnings if existing else (),
            description=str(status.get("description") or (existing.description if existing else "")),
            table_type=str(status.get("table_type") or (existing.table_type if existing else "")),
            location=str(status.get("location") or (existing.location if existing else "")),
            metadata_loaded=True,
        ))
    return catalog


def load_dictionary(
    workbook_path: str | Path,
    table_ids: Mapping[str, str] = TABLE_IDS,
) -> SchemaCatalog:
    """Load optional table and column descriptions; missing tabs are non-fatal."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - exercised in Workbench environment
        raise DictionaryError("Instala openpyxl para leer el diccionario Excel") from exc

    path = Path(workbook_path)
    if not path.is_file():
        raise DictionaryError(f"No existe el diccionario: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: dict[str, TableSchema] = {}
    workbook_warnings: list[str] = []
    try:
        for sheet_name, table_id in table_ids.items():
            if sheet_name not in workbook.sheetnames:
                workbook_warnings.append(f"No existe la hoja `{sheet_name}`; se continuará sin sus descripciones.")
                continue
            rows = workbook[sheet_name].iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                workbook_warnings.append(f"La hoja `{sheet_name}` está vacía; se continuará sin sus descripciones.")
                continue
            headers = {_normalize(value): index for index, value in enumerate(header) if value is not None}
            name_index = headers.get("field name")
            type_index = headers.get("type")
            description_index = headers.get("description")
            if name_index is None:
                workbook_warnings.append(f"La hoja `{sheet_name}` no tiene `field name`; se omitirá.")
                continue

            columns: dict[str, ColumnSchema] = {}
            warnings: list[str] = []
            for row_number, row in enumerate(rows, start=2):
                if not row or name_index >= len(row) or row[name_index] is None:
                    continue
                name = str(row[name_index]).strip()
                if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name):
                    warnings.append(f"Campo inválido en fila {row_number}; omitido.")
                    continue
                normalized_name = name.lower()
                if normalized_name in columns:
                    warnings.append(f"Campo duplicado `{name}` en fila {row_number}; se conserva la primera descripción.")
                    continue
                data_type = (
                    str(row[type_index]).strip().upper()
                    if type_index is not None and type_index < len(row) and row[type_index] is not None
                    else ""
                )
                description = ""
                if description_index is not None and description_index < len(row) and row[description_index] is not None:
                    description = str(row[description_index]).strip()
                    if description.upper() in {"#N/A", "N/A", "NA"}:
                        description = ""
                columns[normalized_name] = ColumnSchema(name, data_type, description)
            tables[_key(table_id)] = TableSchema(
                sheet_name=sheet_name,
                table_id=table_id,
                columns=columns,
                warnings=tuple(warnings),
            )
            workbook_warnings.extend(f"{sheet_name}: {warning}" for warning in warnings)
    finally:
        workbook.close()

    return SchemaCatalog(tables=tables, warnings=tuple(workbook_warnings))
